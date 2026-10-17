"""Money-flow calculations for the Account page (wages, purchases, expenses, funds, budget)."""
from datetime import date, timedelta
from io import BytesIO

import pandas as pd

import db

PERIODS = ["Day", "Week (Sun–Sat)", "Month", "Overall", "Custom"]


# ---------------------------------------------------------------- basics
def inr(x):
    """Format as Indian rupees with lakh/crore grouping, e.g. ₹12,34,567."""
    x = round(float(x or 0))
    neg, s = x < 0, str(abs(x))
    if len(s) > 3:
        head, tail, parts = s[:-3], s[-3:], []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts + [tail])
    return ("-" if neg else "") + "₹" + s


def q(sql, params=()):
    with db.conn() as c:
        return pd.read_sql_query(sql, c, params=params)


def as_dates(df, cols):
    for col in cols:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce")
    return df


def week_end_of(d):
    """Saturday that ends the Sunday–Saturday pay week containing d."""
    return d + timedelta(days=(5 - d.weekday()) % 7)


def period_range(kind, anchor, custom=None):
    """Return (start, end) dates for a period, or (None, None) for Overall."""
    if kind == "Day":
        return anchor, anchor
    if kind == "Week (Sun–Sat)":
        we = week_end_of(anchor)
        return we - timedelta(days=6), we
    if kind == "Month":
        first = anchor.replace(day=1)
        return first, (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    if kind == "Custom" and custom:
        return custom
    return None, None


def in_period(series, start, end):
    """Boolean mask for a datetime Series inside [start, end] (all True when start is None)."""
    if start is None:
        return series.notna() | series.isna()
    return (series >= pd.Timestamp(start)) & (series <= pd.Timestamp(end))


# ---------------------------------------------------------------- generic table sync (edit / add / delete)
def _conv(v, kind):
    if kind == "date":
        return db.iso(v)
    if kind == "num":
        return None if pd.isna(v) else float(v)
    if kind == "int":
        return None if pd.isna(v) else int(v)
    return db.clean_text(v)


def sync_table(table, spec, edited, shown_ids):
    """Write an edited DataFrame (with an 'id' column) back to `table`.
    spec = [(column, kind, required)], kind in date/num/int/text. Rows missing a required value are left unchanged."""
    names = [s[0] for s in spec]
    kept = set()
    with db.conn() as c:
        for _, r in edited.iterrows():
            rid = None if ("id" not in r or pd.isna(r["id"])) else int(r["id"])
            if rid is not None:
                kept.add(rid)
            vals = [_conv(r[n], k) for n, k, _ in spec]
            if any(req and (v is None or v == "") for (_, _, req), v in zip(spec, vals)):
                continue
            if rid is None:
                c.execute(f"INSERT INTO {table} ({','.join(names)}) VALUES ({','.join('?' * len(names))})", vals)
            else:
                c.execute(f"UPDATE {table} SET {','.join(n + '=?' for n in names)} WHERE id=?", vals + [rid])
        for rid in set(shown_ids) - kept:
            c.execute(f"DELETE FROM {table} WHERE id=?", (rid,))


# ---------------------------------------------------------------- wages
def workers_all():
    return q("SELECT id, name, role, agency, active, COALESCE(daily_wage, 0) AS daily_wage FROM workers ORDER BY name")


def wages_week(week_end):
    """Weekly wage sheet for the Sunday–Saturday week ending `week_end` (a Saturday)."""
    ws = week_end - timedelta(days=6)
    att = q("SELECT w.id AS worker_id, w.name AS Worker, w.role AS Role, w.agency AS Agency, "
            "COALESCE(w.daily_wage, 0) AS rate, "
            "SUM(CASE a.status WHEN 'Present' THEN 1 ELSE 0 END) AS present, "
            "SUM(CASE a.status WHEN 'Half day' THEN 1 ELSE 0 END) AS half "
            "FROM attendance a JOIN workers w ON w.id = a.worker_id "
            "WHERE a.att_date BETWEEN ? AND ? GROUP BY w.id ORDER BY w.name", (ws.isoformat(), week_end.isoformat()))
    paid = q("SELECT * FROM wage_payments WHERE week_end = ?", (week_end.isoformat(),)).set_index("worker_id")
    adv = q("SELECT worker_id, SUM(amount) AS v FROM advances WHERE adv_date <= ? GROUP BY worker_id",
            (week_end.isoformat(),)).set_index("worker_id")["v"]
    ded = q("SELECT worker_id, SUM(deduction) AS v FROM wage_payments WHERE week_end < ? GROUP BY worker_id",
            (week_end.isoformat(),)).set_index("worker_id")["v"]
    rows = []
    for _, r in att.iterrows():
        wid = int(r["worker_id"])
        days = float(r["present"]) + 0.5 * float(r["half"])
        balance = float(adv.get(wid, 0)) - float(ded.get(wid, 0))
        if wid in paid.index:
            p = paid.loc[wid]
            rate, gross, deduction, status, paid_on = p["rate"], p["gross"], p["deduction"], "Paid", p["paid_on"]
            days = p["days"]
        else:
            rate = float(r["rate"])
            gross, deduction, status, paid_on = days * rate, 0.0, "Pending", None
        rows.append({
            "ID": wid, "Worker": r["Worker"], "Role": r["Role"], "Agency": r["Agency"],
            "Present": int(r["present"]), "Half day": int(r["half"]), "Days payable": days,
            "Daily wage": float(rate), "Gross wages": float(gross),
            "Advance balance": balance, "Deduct advance": float(deduction),
            "Net payable": float(gross) - float(deduction), "Status": status, "Paid on": paid_on,
            "Pay": status == "Pending" and float(rate) > 0,
        })
    cols = ["ID", "Worker", "Role", "Agency", "Present", "Half day", "Days payable", "Daily wage", "Gross wages",
            "Advance balance", "Deduct advance", "Net payable", "Status", "Paid on", "Pay"]
    df = pd.DataFrame(rows, columns=cols)
    df["Paid on"] = pd.to_datetime(df["Paid on"], errors="coerce")
    return df


def week_attendance_grid(week_end):
    """Worker x day grid (P / ½ / A / L) for the week, to cross-check before paying."""
    ws = week_end - timedelta(days=6)
    a = q("SELECT w.name AS Worker, a.att_date AS d, a.status AS s FROM attendance a JOIN workers w ON w.id = a.worker_id "
          "WHERE a.att_date BETWEEN ? AND ?", (ws.isoformat(), week_end.isoformat()))
    if a.empty:
        return pd.DataFrame()
    a["day"] = pd.to_datetime(a["d"]).dt.strftime("%a %d")
    a["s"] = a["s"].map({"Present": "P", "Half day": "½", "Absent": "A", "Leave": "L"})
    order = [(ws + timedelta(days=i)).strftime("%a %d") for i in range(7)]
    g = a.pivot_table(index="Worker", columns="day", values="s", aggfunc="first").reindex(columns=order)
    return g.fillna("–").reset_index()


def record_wage_payments(week_end, edited, paid_on):
    """Save payments for ticked, still-pending workers. Returns (saved_count, error list)."""
    errors, todo = [], []
    for _, r in edited.iterrows():
        if r["Status"] != "Pending" or not bool(r["Pay"]):
            continue
        gross = float(r["Gross wages"])
        if float(r["Daily wage"]) <= 0 and float(r["Days payable"]) > 0:
            errors.append(f"{r['Worker']}: daily wage is not set – set it in the Rates & Advances tab first.")
            continue
        ded = 0.0 if pd.isna(r["Deduct advance"]) else float(r["Deduct advance"])
        if ded < 0:
            errors.append(f"{r['Worker']}: deduction cannot be negative.")
        elif ded > float(r["Advance balance"]) + 1e-9:
            errors.append(f"{r['Worker']}: deduction {inr(ded)} is more than the advance balance {inr(r['Advance balance'])}.")
        elif ded > gross + 1e-9:
            errors.append(f"{r['Worker']}: deduction {inr(ded)} is more than the gross wages {inr(gross)}.")
        else:
            todo.append((int(r["ID"]), float(r["Days payable"]), float(r["Daily wage"]), gross, ded))
    if errors:
        return 0, errors
    with db.conn() as c:
        for wid, days, rate, gross, ded in todo:
            c.execute("INSERT OR REPLACE INTO wage_payments (week_end, worker_id, days, rate, gross, deduction, net, paid_on) "
                      "VALUES (?,?,?,?,?,?,?,?)",
                      (week_end.isoformat(), wid, days, rate, gross, ded, gross - ded, paid_on.isoformat()))
    return len(todo), []


def cancel_wage_payments(week_end):
    with db.conn() as c:
        c.execute("DELETE FROM wage_payments WHERE week_end = ?", (week_end.isoformat(),))


def wages_by_week():
    """One row per Sunday–Saturday week that has attendance: man-days, gross wages, paid status."""
    a = q("SELECT a.att_date AS d, a.status AS s, COALESCE(w.daily_wage, 0) AS rate FROM attendance a "
          "JOIN workers w ON w.id = a.worker_id WHERE a.status IN ('Present', 'Half day')")
    if a.empty:
        return pd.DataFrame(columns=["Week ending (Sat)", "Man-days", "Gross wages", "Paid (net)", "Advance recovered", "Status"])
    a["d"] = pd.to_datetime(a["d"]).dt.date
    a["we"] = a["d"].map(week_end_of)
    a["days"] = a["s"].map({"Present": 1.0, "Half day": 0.5})
    a["amt"] = a["days"] * a["rate"]
    g = a.groupby("we").agg(**{"Man-days": ("days", "sum"), "Gross wages": ("amt", "sum")}).reset_index()
    p = q("SELECT week_end, SUM(gross) AS pg, SUM(net) AS pn, SUM(deduction) AS pd FROM wage_payments GROUP BY week_end")
    p["we"] = pd.to_datetime(p["week_end"]).dt.date
    g = g.merge(p[["we", "pg", "pn", "pd"]], on="we", how="left").fillna({"pg": 0, "pn": 0, "pd": 0})
    g["Status"] = g.apply(lambda r: "Paid" if r["pg"] > 0 and r["pg"] >= r["Gross wages"] - 0.5
                          else ("Part paid" if r["pg"] > 0 else "Pending"), axis=1)
    g = g.rename(columns={"we": "Week ending (Sat)", "pn": "Paid (net)", "pd": "Advance recovered"})
    return g[["Week ending (Sat)", "Man-days", "Gross wages", "Paid (net)", "Advance recovered", "Status"]].sort_values(
        "Week ending (Sat)", ascending=False)


def advance_summary():
    w = workers_all()
    adv = q("SELECT worker_id, SUM(amount) AS given FROM advances GROUP BY worker_id").set_index("worker_id")["given"]
    ded = q("SELECT worker_id, SUM(deduction) AS rec FROM wage_payments GROUP BY worker_id").set_index("worker_id")["rec"]
    w["Advance given"] = w["id"].map(adv).fillna(0.0)
    w["Recovered"] = w["id"].map(ded).fillna(0.0)
    w["Balance"] = w["Advance given"] - w["Recovered"]
    out = w[w["Advance given"] > 0][["name", "role", "Advance given", "Recovered", "Balance"]]
    return out.rename(columns={"name": "Worker", "role": "Role"})


# ---------------------------------------------------------------- purchases / expenses / funds / supplier payments
def purchases_df():
    df = q("SELECT * FROM purchases")
    df["amount"] = df["amount"].where(df["amount"].notna(), df["quantity"] * df["rate"]).fillna(0.0)
    df = as_dates(df, ["order_date", "expected_date", "received_date"])
    return df


def purchases_active():
    df = purchases_df()
    return df[df["status"] != "Cancelled"]


def purchase_head(cat):
    return db.PURCHASE_HEAD.get(cat, "Materials")


def expenses_df():
    return as_dates(q("SELECT * FROM expenses ORDER BY exp_date DESC, id DESC"), ["exp_date"])


def funds_df():
    return as_dates(q("SELECT * FROM fund_receipts ORDER BY rec_date DESC, id DESC"), ["rec_date"])


def supplier_payments_df():
    return as_dates(q("SELECT * FROM supplier_payments ORDER BY pay_date DESC, id DESC"), ["pay_date"])


def supplier_ledger():
    p = purchases_active()
    p = p.assign(supplier=p["supplier"].replace("", "(no supplier)"))
    bought = p.groupby("supplier")["amount"].sum()
    pay = supplier_payments_df()
    pay = pay.assign(supplier=pay["supplier"].replace("", "(no supplier)"))
    paid = pay.groupby("supplier")["amount"].sum()
    led = pd.DataFrame({"Purchased": bought, "Paid": paid}).fillna(0.0)
    led["Outstanding"] = led["Purchased"] - led["Paid"]
    led.index.name = "Supplier"
    return led.reset_index().sort_values("Outstanding", ascending=False)


# ---------------------------------------------------------------- totals, heads, budget
def wages_accrued_total():
    r = q("SELECT COALESCE(SUM(CASE a.status WHEN 'Present' THEN 1.0 WHEN 'Half day' THEN 0.5 ELSE 0 END "
          "* COALESCE(w.daily_wage, 0)), 0) AS v FROM attendance a JOIN workers w ON w.id = a.worker_id")
    return float(r["v"].iloc[0])


def cost_by_head():
    """Cost incurred per head: wages from attendance, purchases and other expenses."""
    rows = {h: 0.0 for h in db.HEADS}
    rows["Labour (wages)"] = wages_accrued_total()
    for _, r in purchases_active().iterrows():
        rows[purchase_head(r["category"])] += float(r["amount"])
    for _, r in expenses_df().iterrows():
        h = r["head"] if r["head"] in rows else "Miscellaneous"
        rows[h] += float(r["amount"] or 0)
    return rows


def budget_table():
    b = q("SELECT head, amount FROM budgets").set_index("head")["amount"]
    actual = cost_by_head()
    df = pd.DataFrame({"Head": db.HEADS})
    df["Budget"] = df["Head"].map(b).fillna(0.0)
    df["Actual"] = df["Head"].map(actual).fillna(0.0)
    df["Balance"] = df["Budget"] - df["Actual"]
    df["% used"] = (df["Actual"] / df["Budget"].where(df["Budget"] > 0)).mul(100).round(1)
    return df


def save_budget(df):
    with db.conn() as c:
        for _, r in df.iterrows():
            c.execute("INSERT OR REPLACE INTO budgets (head, amount) VALUES (?, ?)",
                      (r["Head"], 0.0 if pd.isna(r["Budget"]) else float(r["Budget"])))


def cash_summary():
    """Money in, money out (cash basis), and what is still to be paid."""
    funds = float(q("SELECT COALESCE(SUM(amount),0) AS v FROM fund_receipts")["v"].iloc[0])
    adv = float(q("SELECT COALESCE(SUM(amount),0) AS v FROM advances")["v"].iloc[0])
    wnet = float(q("SELECT COALESCE(SUM(net),0) AS v FROM wage_payments")["v"].iloc[0])
    wgross_paid = float(q("SELECT COALESCE(SUM(gross),0) AS v FROM wage_payments")["v"].iloc[0])
    wded = float(q("SELECT COALESCE(SUM(deduction),0) AS v FROM wage_payments")["v"].iloc[0])
    sup = float(q("SELECT COALESCE(SUM(amount),0) AS v FROM supplier_payments")["v"].iloc[0])
    oth = float(q("SELECT COALESCE(SUM(amount),0) AS v FROM expenses")["v"].iloc[0])
    purchased = float(purchases_active()["amount"].sum())
    accrued = wages_accrued_total()
    out = adv + wnet + sup + oth
    return {
        "funds": funds, "labour_cash": adv + wnet, "supplier_cash": sup, "other_cash": oth,
        "cash_out": out, "cash_in_hand": funds - out,
        "wages_accrued": accrued, "wages_pending": max(accrued - wgross_paid, 0.0),
        "advance_balance": adv - wded, "purchased": purchased, "supplier_payable": purchased - sup,
        "cost_incurred": accrued + purchased + oth,
    }


def monthly_cost():
    """Cost incurred per month, split into wages / materials & purchases / other expenses."""
    w = q("SELECT a.att_date AS d, CASE a.status WHEN 'Present' THEN 1.0 WHEN 'Half day' THEN 0.5 ELSE 0 END "
          "* COALESCE(w.daily_wage, 0) AS amt FROM attendance a JOIN workers w ON w.id = a.worker_id")
    parts = {}
    if not w.empty:
        w["m"] = pd.to_datetime(w["d"]).dt.strftime("%Y-%m")
        parts["Wages"] = w.groupby("m")["amt"].sum()
    p = purchases_active().dropna(subset=["order_date"])
    if not p.empty:
        parts["Purchases"] = p.groupby(p["order_date"].dt.strftime("%Y-%m"))["amount"].sum()
    e = expenses_df().dropna(subset=["exp_date"])
    if not e.empty:
        parts["Other expenses"] = e.groupby(e["exp_date"].dt.strftime("%Y-%m"))["amount"].sum()
    return pd.DataFrame(parts).fillna(0.0).sort_index() if parts else pd.DataFrame()


def monthly_cashflow():
    parts = {}
    f = funds_df().dropna(subset=["rec_date"])
    if not f.empty:
        parts["Funds received"] = f.groupby(f["rec_date"].dt.strftime("%Y-%m"))["amount"].sum()
    out = {}
    adv = as_dates(q("SELECT adv_date AS d, amount FROM advances"), ["d"]).dropna(subset=["d"])
    wp = as_dates(q("SELECT paid_on AS d, net AS amount FROM wage_payments"), ["d"]).dropna(subset=["d"])
    sp = supplier_payments_df().rename(columns={"pay_date": "d"})[["d", "amount"]].dropna(subset=["d"])
    ex = expenses_df().rename(columns={"exp_date": "d"})[["d", "amount"]].dropna(subset=["d"])
    allout = pd.concat([adv, wp, sp, ex])
    if not allout.empty:
        parts["Money paid out"] = allout.groupby(allout["d"].dt.strftime("%Y-%m"))["amount"].sum()
    return pd.DataFrame(parts).fillna(0.0).sort_index() if parts else pd.DataFrame()


def floor_progress():
    with db.conn() as c:
        rows = [dict(r) for r in c.execute("SELECT * FROM activities").fetchall()]
    today = date.today()
    out = []
    for floor in db.FLOORS:
        total = done = delayed = 0
        for r in (x for x in rows if x["floor"] == floor):
            ps, pe = db.to_date(r["planned_start"]), db.to_date(r["planned_end"])
            a_s, a_e = db.to_date(r["actual_start"]), db.to_date(r["actual_end"])
            st_ = r["status"] if r["status"] in db.STATUS_OPTIONS else db.derive_status(a_s, a_e)
            st_ = db.sync_status(st_, a_s, a_e)
            total += 1
            done += st_ == "Completed"
            delayed += db.delay_days(ps, pe, a_s, a_e, st_, today) > 0
        out.append({"Floor": floor, "Activities": total, "Completed": done, "Delayed": delayed,
                    "% complete": round(100 * done / total) if total else 0})
    return pd.DataFrame(out)


def schedule_summary():
    """Progress and delay figures taken from the Schedule page."""
    with db.conn() as c:
        rows = [dict(r) for r in c.execute("SELECT * FROM activities").fetchall()]
    today = date.today()
    total, done, delayed, delay_days_total = len(rows), 0, 0, 0
    next_concrete = None
    for r in rows:
        ps, pe = db.to_date(r["planned_start"]), db.to_date(r["planned_end"])
        a_s, a_e = db.to_date(r["actual_start"]), db.to_date(r["actual_end"])
        status = r["status"] if r["status"] in db.STATUS_OPTIONS else db.derive_status(a_s, a_e)
        status = db.sync_status(status, a_s, a_e)
        if status == "Completed":
            done += 1
        d = db.delay_days(ps, pe, a_s, a_e, status, today)
        if d > 0:
            delayed += 1
            delay_days_total += d
        if r["is_concrete"] and status != "Completed" and ps >= today:
            if next_concrete is None or ps < next_concrete[1]:
                next_concrete = (f'{r["floor"]} – {r["name"]}', ps)
    return {"total": total, "done": done, "delayed": delayed, "delay_days": delay_days_total,
            "next_concrete": next_concrete, "progress": (done / total if total else 0.0)}


def insights(cs, ss, budget):
    """Plain-language observations a site engineer would want to see."""
    out = []
    tb = float(budget["Budget"].sum())
    if tb > 0:
        spent = cs["cost_incurred"] / tb
        out.append(f"Cost incurred is **{spent:.0%}** of the total budget while **{ss['progress']:.0%}** of activities are complete.")
        if spent > ss["progress"] + 0.10:
            out.append("⚠️ Spending is running **ahead of progress** – check wastage and rates.")
        elif spent < ss["progress"] - 0.10:
            out.append("✅ Spending is **behind progress** – check that all bills are being booked.")
        over = budget[(budget["Budget"] > 0) & (budget["Actual"] > budget["Budget"])]
        for _, r in over.iterrows():
            out.append(f"⚠️ **{r['Head']}** has crossed its budget by {inr(r['Actual'] - r['Budget'])}.")
    else:
        out.append("Set the budget in the **Budget** tab to compare cost with progress.")
    if cs["cash_in_hand"] < 0:
        out.append(f"⚠️ Cash is **short by {inr(-cs['cash_in_hand'])}** – paid out more than funds received.")
    if cs["supplier_payable"] > 0:
        out.append(f"Suppliers are owed **{inr(cs['supplier_payable'])}**.")
    if cs["wages_pending"] > 0:
        out.append(f"Wages pending (not yet recorded as paid): **{inr(cs['wages_pending'])}**.")
    if cs["advance_balance"] > 0:
        out.append(f"Advances to recover from workers: **{inr(cs['advance_balance'])}**.")
    return out


def schedule_insights():
    """Schedule-related observations for the Schedule page (no money figures)."""
    with db.conn() as c:
        rows = [dict(r) for r in c.execute("SELECT * FROM activities ORDER BY floor, seq").fetchall()]
    today = date.today()
    total = done = 0
    delayed, today_list, concrete = [], [], []
    for r in rows:
        ps, pe = db.to_date(r["planned_start"]), db.to_date(r["planned_end"])
        a_s, a_e = db.to_date(r["actual_start"]), db.to_date(r["actual_end"])
        st_ = r["status"] if r["status"] in db.STATUS_OPTIONS else db.derive_status(a_s, a_e)
        st_ = db.sync_status(st_, a_s, a_e)
        total += 1
        label = f'{r["floor"]} – {r["name"]}'
        if st_ == "Completed":
            done += 1
            continue
        d = db.delay_days(ps, pe, a_s, a_e, st_, today)
        if d > 0 and st_ != "On hold":
            delayed.append((d, label))
        if ps <= today <= pe and st_ != "On hold":
            today_list.append(label)
        if r["is_concrete"] and ps >= today:
            concrete.append((ps, label))
    out = [f"**{done}** of **{total}** activities are completed ({done / total:.0%})." if total else "No activities."]
    if concrete:
        d0, name = min(concrete)
        n = (d0 - today).days
        when = "today" if n == 0 else f"in {n} day{'s' if n != 1 else ''}"
        out.append(f"Next concrete pour: **{name}** on {d0:%a %d %b %Y} ({when}) – arrange cement, steel, "
                   "mixer / ready-mix and labour in advance.")
    if today_list:
        out.append("Planned for today: " + ", ".join(today_list[:4]) + ("…" if len(today_list) > 4 else "") + ".")
    if delayed:
        delayed.sort(reverse=True)
        out.append(f"⚠️ **{len(delayed)}** activities are delayed – longest: "
                   + ", ".join(f"{n} ({d} day{'s' if d != 1 else ''})" for d, n in delayed[:3]) + ".")
    else:
        out.append("✅ No delayed activities.")
    return out


# ---------------------------------------------------------------- Excel report
def build_report(start, end, week_end):
    """Excel workbook with every money report. start/end = None means the whole project."""
    from openpyxl.styles import Alignment, Font, PatternFill

    def cut(df, col):
        return df[in_period(df[col], start, end)] if start else df

    cs = cash_summary()
    bt = budget_table()
    summary = pd.DataFrame([
        ("Report period", "Overall (whole project)" if start is None else f"{start:%d %b %Y} to {end:%d %b %Y}"),
        ("Generated on", f"{date.today():%d %b %Y}"),
        ("Total budget", inr(bt["Budget"].sum())),
        ("Cost incurred (wages + purchases + other)", inr(cs["cost_incurred"])),
        ("Funds received", inr(cs["funds"])),
        ("Cash paid out", inr(cs["cash_out"])),
        ("Cash in hand / (shortfall)", inr(cs["cash_in_hand"])),
        ("Supplier payable", inr(cs["supplier_payable"])),
        ("Wages pending", inr(cs["wages_pending"])),
        ("Advances to recover", inr(cs["advance_balance"])),
    ], columns=["Item", "Value"])

    pur = cut(purchases_df(), "order_date").drop(columns=["id"]).sort_values("order_date")
    exp = cut(expenses_df(), "exp_date").drop(columns=["id"]).sort_values("exp_date")
    fun = cut(funds_df(), "rec_date").drop(columns=["id"]).sort_values("rec_date")
    spy = cut(supplier_payments_df(), "pay_date").drop(columns=["id"]).sort_values("pay_date")
    wk = wages_week(week_end).drop(columns=["ID", "Pay"])
    att = as_dates(q("SELECT a.att_date AS Date, w.name AS Worker, w.role AS Role, w.agency AS Agency, a.status AS Status, "
                     "a.remarks AS Remarks FROM attendance a "
                     "JOIN workers w ON w.id = a.worker_id ORDER BY a.att_date, w.name"), ["Date"])
    att = cut(att, "Date")
    sheets = {
        "Summary": summary,
        f"Wages week {week_end:%d-%b-%Y}": wk,
        "Wages by week": wages_by_week(),
        "Procurement": pur, "Other expenses": exp, "Funds received": fun,
        "Supplier payments": spy, "Supplier ledger": supplier_ledger(),
        "Budget vs actual": bt, "Advances": advance_summary(), "Attendance": att,
    }
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY", date_format="DD/MM/YYYY") as xw:
        for name, df in sheets.items():
            df = df.copy()
            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    df[col] = df[col].dt.date
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width, 10), 45)
            ws.freeze_panes = "A2"
    return buf.getvalue()